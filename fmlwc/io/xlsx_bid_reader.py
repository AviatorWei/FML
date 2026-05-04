"""XLSX → RawBid adapter.

Reads one manager's sealed-bid submission from the standard FME xlsx template
and converts it into domain-layer ``RawBid`` objects.

Expected filename convention
-----------------------------
    FME_<year>_Bid<N>_<CODE>.xlsx
    e.g.  FME_2024_Bid1_ENG.xlsx

The manager code (``ENG``, ``GER``, …) is extracted from the filename.
Because the xlsx file contains no manager_id integer, the caller is
responsible for resolving code → id and passing it to
``XlsxSubmission.to_raw_bids(manager_id)``.

Sheet structure (Sheet1 / active sheet)
----------------------------------------
    Row 1:  Order | Price | ID | Name | Team | Pos | <metadata cols G–J>
    Row 2+: data rows

A row represents a bid when ``Order`` is a **non-zero integer**.
``Order`` becomes ``rank_in_position``; ``Price`` becomes ``amount`` (M EUR);
``ID`` becomes ``player_id``.  ``Name``, ``Team``, and ``Pos`` are
informational and are stored but not used in domain validation.

Conditional-release bids (negative rank) are not present in the current
template, but the reader passes negative ``Order`` values through unchanged
so future templates can include them.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

from ..domain.auction.bids import RawBid


# ---------------------------------------------------------------------------
# Public data structures
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class BidRow:
    """A single parsed row — player identity plus bid parameters."""
    player_id: int
    amount: int             # million EUR
    rank_in_position: int   # >0 acquisition; <0 conditional release
    # informational fields (not used by domain logic)
    player_name: str | None = None
    real_team: str | None = None
    position_str: str | None = None


@dataclass
class XlsxSubmission:
    """Parsed content of a single manager's bid xlsx file.

    Attributes
    ----------
    manager_code:
        Short code extracted from the filename, e.g. ``"ENG"``.
    source_file:
        Basename of the source file, e.g. ``"FME_2024_Bid1_ENG.xlsx"``.
    round_index:
        Auction-round number parsed from the filename (1-based).
    rows:
        Ordered list of bid rows (only rows where Order != 0).
    """
    manager_code: str
    source_file: str
    round_index: int
    rows: list[BidRow] = field(default_factory=list)

    def to_raw_bids(self, manager_id: int) -> list[RawBid]:
        """Bind a resolved *manager_id* to produce ``RawBid`` objects.

        Parameters
        ----------
        manager_id:
            The integer primary key of this manager in the database.
            The caller is responsible for resolving ``manager_code``
            to the correct id before calling this method.
        """
        return [
            RawBid(
                manager_id=manager_id,
                player_id=row.player_id,
                amount=row.amount,
                rank_in_position=row.rank_in_position,
            )
            for row in self.rows
        ]


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------

class XlsxBidParseError(ValueError):
    """Raised when an xlsx file cannot be parsed as a bid submission."""


# ---------------------------------------------------------------------------
# Defaults
# ---------------------------------------------------------------------------

# FME_2024_Bid1_ENG.xlsx  →  group(1)=round_index, group(2)=manager_code
DEFAULT_FILENAME_PATTERN: re.Pattern = re.compile(
    r'^FME_\d{4}_Bid(\d+)_([A-Z]{2,4})\.xlsx$',
    re.IGNORECASE,
)

DEFAULT_HEADERS: tuple[str, ...] = ("Order", "Price", "ID", "Name", "Team", "Pos")


# ---------------------------------------------------------------------------
# Reader
# ---------------------------------------------------------------------------

class XlsxBidReader:
    """Reads bid xlsx files from a configured directory.

    All format details — the base directory, filename pattern, and expected
    sheet headers — are supplied once at construction time so call sites only
    need to name the file (or call ``read_all``).

    Parameters
    ----------
    directory:
        Folder that contains (or will contain) the xlsx bid files.
    filename_pattern:
        Compiled ``re.Pattern`` (or a plain string that will be compiled with
        ``re.IGNORECASE``) used to recognise valid filenames and extract
        ``round_index`` from group 1 and ``manager_code`` from group 2.
        Defaults to ``DEFAULT_FILENAME_PATTERN``.
    expected_headers:
        Ordered tuple of column header strings expected in row 1 of the
        active sheet.  Defaults to ``DEFAULT_HEADERS``.

    Typical usage::

        reader = XlsxBidReader("example/bids-1/")

        # Read every matching file in the directory
        for sub in reader.read_all():
            manager_id = resolve(sub.manager_code)
            auction_service.submit(round_id, manager_id,
                                   sub.to_raw_bids(manager_id), received_at=now)

        # Read a single file by name
        sub = reader.read_file("FME_2024_Bid1_ENG.xlsx")
    """

    def __init__(
        self,
        directory: Path | str,
        *,
        filename_pattern: re.Pattern | str = DEFAULT_FILENAME_PATTERN,
        expected_headers: tuple[str, ...] = DEFAULT_HEADERS,
    ) -> None:
        self.directory = Path(directory)
        self.filename_pattern: re.Pattern = (
            re.compile(filename_pattern, re.IGNORECASE)
            if isinstance(filename_pattern, str)
            else filename_pattern
        )
        self.expected_headers = expected_headers

    # -- public API --------------------------------------------------------

    def read_file(self, filename: str) -> XlsxSubmission:
        """Parse *filename* (relative to ``self.directory``).

        Parameters
        ----------
        filename:
            Bare filename, e.g. ``"FME_2024_Bid1_ENG.xlsx"``.

        Raises
        ------
        XlsxBidParseError
            If the filename does not match the configured pattern or the
            sheet structure is invalid.
        FileNotFoundError
            If the file does not exist inside ``self.directory``.
        ImportError
            If *openpyxl* is not installed.
        """
        path = self.directory / filename
        manager_code, round_index = self._parse_filename(filename)
        rows = list(self._iter_rows(path))
        return XlsxSubmission(
            manager_code=manager_code,
            source_file=filename,
            round_index=round_index,
            rows=rows,
        )

    def read_all(self) -> list[XlsxSubmission]:
        """Parse every file in ``self.directory`` that matches the filename pattern.

        Non-matching files are silently skipped.
        Results are sorted by ``(round_index, manager_code)``.

        Raises
        ------
        XlsxBidParseError
            On the first file that fails structural validation.
        """
        submissions = []
        for p in sorted(self.directory.glob("*.xlsx")):
            if not self.filename_pattern.match(p.name):
                continue
            submissions.append(self.read_file(p.name))
        submissions.sort(key=lambda s: (s.round_index, s.manager_code))
        return submissions

    # -- internals ---------------------------------------------------------

    def _parse_filename(self, name: str) -> tuple[str, int]:
        """Return ``(manager_code, round_index)`` or raise ``XlsxBidParseError``."""
        m = self.filename_pattern.match(name)
        if not m:
            raise XlsxBidParseError(
                f"Filename {name!r} does not match the configured pattern "
                f"{self.filename_pattern.pattern!r}"
            )
        return m.group(2).upper(), int(m.group(1))

    def _iter_rows(self, path: Path) -> Iterator[BidRow]:
        """Yield one ``BidRow`` per non-zero-order row in the active sheet."""
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportError(
                "openpyxl is required to read xlsx files: pip install openpyxl"
            ) from exc

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb.active
            row_iter = ws.iter_rows(values_only=True)

            # --- validate header ---
            try:
                header = next(row_iter)
            except StopIteration:
                raise XlsxBidParseError(f"{path.name}: file has no rows")

            for col_idx, expected in enumerate(self.expected_headers):
                actual = header[col_idx] if col_idx < len(header) else None
                if str(actual).strip() != expected:
                    raise XlsxBidParseError(
                        f"{path.name}: column {col_idx + 1} header is "
                        f"{actual!r}, expected {expected!r}"
                    )

            # --- data rows ---
            for row_num, row in enumerate(row_iter, start=2):
                if len(row) < 3:
                    continue

                order = row[0]
                price = row[1]
                pid   = row[2]
                name  = row[3] if len(row) > 3 else None
                team  = row[4] if len(row) > 4 else None
                pos   = row[5] if len(row) > 5 else None

                # Rows with no bid: Order is None or 0
                if order is None or order == 0:
                    continue

                # --- type checks ---
                if not isinstance(order, (int, float)) or isinstance(order, bool):
                    raise XlsxBidParseError(
                        f"{path.name} row {row_num}: "
                        f"Order must be numeric, got {order!r}"
                    )
                if not isinstance(price, (int, float)) or isinstance(price, bool):
                    raise XlsxBidParseError(
                        f"{path.name} row {row_num}: "
                        f"Price must be numeric, got {price!r}"
                    )
                if not isinstance(pid, (int, float)) or isinstance(pid, bool):
                    raise XlsxBidParseError(
                        f"{path.name} row {row_num}: "
                        f"ID must be numeric, got {pid!r}"
                    )

                yield BidRow(
                    player_id=int(pid),
                    amount=int(price),
                    rank_in_position=int(order),
                    player_name=str(name) if name is not None else None,
                    real_team=str(team) if team is not None else None,
                    position_str=str(pos) if pos is not None else None,
                )
        finally:
            wb.close()
